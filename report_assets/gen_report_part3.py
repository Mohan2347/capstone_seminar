"""Part 3: Chapters 5-9"""
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUTPUT = r'C:\Users\HP\Downloads\InternAI_Compass_Capstone_Report.docx'
doc = Document(OUTPUT)

def sf(run, size=12, bold=False, color=None):
    run.font.name = 'Times New Roman'; run.font.size = Pt(size); run.font.bold = bold
    if color: run.font.color.rgb = RGBColor(*color)

def h1(t):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(24); p.paragraph_format.space_after = Pt(12)
    r = p.add_run(t); sf(r, 16, True, (0,0,0))

def h2(t):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(12); p.paragraph_format.space_after = Pt(6)
    r = p.add_run(t); sf(r, 13, True, (0,0,0))

def body(t):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY; p.paragraph_format.space_after = Pt(8)
    r = p.add_run(t); sf(r, 12)

def pb(): doc.add_page_break()

# ===== CHAPTER 5 =====
h1('CHAPTER 5: SOFTWARE REQUIREMENT ANALYSIS')
h2('5.1 Introduction')
body('Software requirements for InternAI Compass were derived from three sources: a systematic analysis of limitations identified in existing platforms reviewed in Chapter 3; the capabilities of available AI and web development technologies; and the authentic needs of the two primary user groups \u2014 students seeking personalized internship guidance and administrators requiring institutional-level placement oversight. Requirements are organized into functional, non-functional, and system categories.')

h2('5.2 General Description')
body('InternAI Compass is a browser-based web application accessible from any modern desktop browser without installation or configuration. Users authenticate through a managed identity provider supporting email/password registration and social login. Upon authentication, the system routes users to their respective interface based on assigned role \u2014 student, employer, or administrator.')
body('The core of the platform is the SmartMatch-AI recommendation engine, which operates continuously in the background, generating and refining internship matches based on each student\u2019s evolving profile and interaction history. Every application, view, or feedback interaction updates the reinforcement learning state, causing the recommendation weights to shift toward configurations that better predict user satisfaction.')

h2('5.3 Specific Requirements')
h2('5.3.1 FUNCTIONAL REQUIREMENTS')
body('FR-01: The system shall allow visitors to register accounts through email/password or social login.')
body('FR-02: Authenticated students shall be able to create and update multidimensional profiles including academic records, skills, work experience, preferences, and personality traits.')
body('FR-03: The system shall generate 768-dimensional BERT embeddings for student profile text and internship description text.')
body('FR-04: The system shall compute cosine similarity between student and internship embeddings as the content-based component of the hybrid score.')
body('FR-05: The system shall apply collaborative filtering using an interaction matrix of application and feedback history from similar users.')
body('FR-06: The SmartMatch-AI algorithm shall compute a hybrid score H(s,i) = \u03b1 \u00d7 Simcontent(s,i) + (1\u2212\u03b1) \u00d7 Simcollab(s,i) for each student-internship pair.')
body('FR-07: The reinforcement learning module shall update the hybrid weight parameter \u03b1 using the formula \u03b1t+1 = \u03b1t + \u03b7(rt \u2212 r\u0302t) after each user interaction.')
body('FR-08: Employers shall be required to complete a blockchain verification process before their postings become visible to students.')
body('FR-09: The system shall generate a skill-gap analysis for each student comparing their current skills against the competency requirements of recommended internships.')
body('FR-10: Students shall receive personalized learning path recommendations based on identified skill gaps.')
body('FR-11: Administrators shall have access to analytics dashboards showing application trends, placement rates, and student readiness indicators.')
body('FR-12: Students shall be able to view their full application history, recommendation history, and interaction feedback log.')

h2('5.3.2 NON-FUNCTIONAL REQUIREMENTS')
body('NFR-01 Performance: End-to-end recommendation generation shall complete within three seconds from profile update or interaction trigger.')
body('NFR-02 Scalability: The serverless architecture shall support horizontal scaling without code changes.')
body('NFR-03 Security: Non-public routes shall be protected by role-based access control middleware.')
body('NFR-04 Accuracy: The recommendation engine shall achieve accuracy above 90%, precision above 88%, and adaptability above 87%.')
body('NFR-05 Usability: A first-time user shall be able to complete profile creation and receive initial recommendations within five minutes without instructions.')
body('NFR-06 Reliability: All components shall include fallback defaults to prevent pipeline failure from malformed inputs.')
body('NFR-07 Responsiveness: Reinforcement learning weight updates shall complete within 200 milliseconds of receiving user feedback.')
body('NFR-08 Data Integrity: Composite unique constraints shall prevent duplicate application records.')

h2('5.3.3 SYSTEM REQUIREMENTS')
body('Development and deployment environment:')
body('Node.js 18+ for Next.js 15 runtime; Python 3.10+ for AI microservice; PostgreSQL 14+ on NeonDB serverless; Client access through Chrome 100+, Firefox 100+, or Edge 100+; Outbound internet access for Gemini API, Hugging Face, and blockchain verification services.')
body('Environment variables: DATABASE_URL, GEMINI_API_KEY, CLERK_SECRET_KEY, CLERK_PUBLISHABLE_KEY, WEBHOOK_SECRET, and ADMIN_EMAIL.')

pb()

# ===== CHAPTER 6 =====
h1('CHAPTER 6: DESIGN')
h2('6.1 System Design')
body('The system architecture follows a three-layer microservices pattern. The Presentation Layer consists of the Next.js web application serving both the student-facing frontend and the administrative dashboard through server-side rendering and API routes. The Application Layer handles authentication, business logic, and orchestration through API routes that communicate with both the database and the AI microservice. The AI Engine Layer is a dedicated Python FastAPI microservice responsible for embedding generation, similarity computation, collaborative filtering, reinforcement learning updates, and skill-gap analysis.')
body('Role-based access control divides the application into three interface zones: public routes accessible without authentication; student routes for profile management and recommendation browsing; and administrator routes for institutional management and analytics.')

h2('6.2 Database Design')
body('The relational database schema comprises four primary models: Student, Internship, Application, and Interaction. The Student model stores profile data including academic records, skills vectors, preferences, and role assignments. The Internship model holds posting details along with pre-computed BERT embedding vectors for efficient similarity computation. The Application model records each student-internship pairing with timestamps and status. The Interaction model captures behavioral signals \u2014 views, saves, and applications \u2014 used as feedback by the reinforcement learning module.')

h2('6.3 Design Notations')
body('The application follows consistent architectural conventions. All TypeScript interfaces and enumerations are defined in src/types/index.ts. Database operations use a singleton Prisma client in src/lib/prisma.ts to prevent connection pool exhaustion in serverless deployments. The embedding generation and similarity computation functions are implemented in src/lib/smartmatch.ts. Components are organized into domain folders within src/components/.')

h2('6.4 Detailed Design')
h2('6.4.1 SmartMatch-AI Hybrid Algorithm')
body('The SmartMatch-AI algorithm is the core intelligence engine of InternAI Compass. For each student-internship pair, the algorithm computes a hybrid score that combines semantic content similarity with behavioral collaborative signals.')
body('The content similarity component is computed as the cosine similarity between the student\u2019s BERT embedding vector Vs and the internship\u2019s BERT embedding vector Vi: Simcontent(s,i) = (Vs \u00b7 Vi) / (||Vs|| \u00d7 ||Vi||).')
body('The collaborative component is derived from an interaction matrix Ru,i capturing historical feedback from users with similar profiles: Simcollab(s,i) = \u03a3(Ru,i \u00d7 Rs,i) / \u221a(\u03a3Ru,i\u00b2 \u00d7 \u03a3Rs,i\u00b2).')
body('The final hybrid score is: H(s,i) = \u03b1 \u00d7 Simcontent(s,i) + (1\u2212\u03b1) \u00d7 Simcollab(s,i), where \u03b1 is the adaptive weight parameter maintained by the reinforcement learning module.')

h2('6.4.2 BERT Embedding and Similarity Computation')
body('Student profiles and internship descriptions are transformed into 768-dimensional embedding vectors using a pre-trained BERT transformer model. These embeddings encode contextual semantic meaning, enabling the system to recognize that synonymous skills and role descriptions represent similar competencies even when expressed using different vocabulary.')
body('Embeddings for internship descriptions are pre-computed at posting time and stored in the database to minimize real-time computational load. Student profile embeddings are regenerated whenever the profile is updated, ensuring that recommendations always reflect the student\u2019s current state.')

h2('6.4.3 Reinforcement Learning Feedback Loop')
body('The reinforcement learning mechanism treats each user interaction as a signal for adjusting the hybrid weight parameter \u03b1. When a student views, saves, or applies for a recommended internship, the system records a positive reward signal. When a recommendation is dismissed or an application is rejected, a negative signal is generated.')
body('The weight update rule is: \u03b1t+1 = \u03b1t + \u03b7(rt \u2212 r\u0302t), where \u03b7 is the learning rate, rt is the actual reward observed from the interaction, and r\u0302t is the predicted reward at the time of recommendation.')

h2('6.5 Flowcharts')
h2('6.5.1 STUDENT REGISTRATION AND PROFILE SETUP FLOW')
body('Student navigates to /sign-up and completes registration. Profile creation form collects academic background, skills, work experience, preferences, and personality traits. BERT embedding is generated from profile text and stored. Initial recommendations are computed using SmartMatch-AI with default \u03b1=0.5. Student views personalized recommendation dashboard and begins interaction.')

h2('6.5.2 RECOMMENDATION GENERATION FLOW')
body('Trigger received from profile update or user interaction. Retrieve student embedding vector from database. Retrieve all verified internship embedding vectors. Compute cosine similarity scores for all student-internship pairs. Apply collaborative filtering to incorporate peer interaction signals. Compute hybrid score H(s,i) using current \u03b1. Filter internships above relevance threshold. Rank by hybrid score. Return top-K recommendations to student dashboard. Process user interaction feedback. Update \u03b1 using reinforcement learning weight update formula.')

h2('6.6 Pseudo Code')
h2('6.6.1 SMARTMATCH-AI RECOMMENDATION')
body('FUNCTION SmartMatch(studentProfile, internshipList, alpha, feedbackMatrix):')
body('    Vs = BERT_EMBED(studentProfile)')
body('    scores = []')
body('    FOR each internship i IN internshipList:')
body('        Vi = internship.embedding')
body('        content_sim = COSINE_SIMILARITY(Vs, Vi)')
body('        collab_sim = COLLABORATIVE_SCORE(student, i, feedbackMatrix)')
body('        hybrid = alpha * content_sim + (1 - alpha) * collab_sim')
body('        scores.append((i, hybrid))')
body('    SORT scores BY hybrid DESC')
body('    RETURN TOP_K(scores)')

h2('6.6.2 REINFORCEMENT LEARNING WEIGHT UPDATE')
body('FUNCTION updateAlpha(alpha, eta, actual_reward, predicted_reward):')
body('    return alpha + eta * (actual_reward - predicted_reward)')

pb()

# ===== CHAPTER 7 =====
h1('CHAPTER 7: TESTING')
h2('7.1 Functional Testing')
body('Functional testing was conducted to verify that all specified requirements are correctly implemented. Testing covered the complete student workflow from registration through profile creation, recommendation generation, skill-gap analysis, and application submission. The employer registration and blockchain verification workflow was tested independently. The administrator analytics dashboard was verified against known datasets.')
body('Authentication was confirmed through account registration using both email/password and social login, with correct role assignment verified for student, employer, and administrator accounts. Profile creation was tested with complete and partial data to confirm appropriate handling of incomplete inputs. Embedding generation was verified by comparing output vectors for semantically similar and dissimilar profile texts and confirming expected similarity relationships.')

h2('7.2 Structural Testing')
body('Structural testing examined internal code paths, data transformations, and API response handling. API routes were tested for correct validation responses including 400 for invalid inputs, 401 for unauthenticated requests, and 403 for unauthorized role access. The BERT embedding pipeline was tested with edge-case inputs including extremely short profile descriptions, profiles containing only numeric data, and descriptions with unconventional formatting.')
body('The cosine similarity computation was verified against manually calculated expected values for known vector pairs. The collaborative filtering module was tested with sparse and dense interaction matrices to confirm consistent output. The reinforcement learning weight update was verified to converge toward expected optimal values across extended interaction sequences.')

h2('7.3 Levels of Testing')
body('Unit Testing: Individual functions including BERT_EMBED(), COSINE_SIMILARITY(), COLLABORATIVE_SCORE(), and updateAlpha() were tested in isolation using curated input-output datasets.')
body('Integration Testing: The complete recommendation pipeline was tested end-to-end, verifying correct data flow from profile submission through embedding generation, similarity computation, hybrid scoring, ranking, and dashboard display.')
body('System Testing: The full deployed system on Vercel was tested with complete environment configuration to verify all routes, authentication flows, and API integrations function correctly in production.')
body('User Acceptance Testing: Complete student and administrator workflows were executed from fresh account creation through all major features to confirm the system meets usability and functionality requirements.')

h2('7.4 Testing the Project \u2014 Results Summary')
# Testing results table
table = doc.add_table(rows=6, cols=4)
table.style = 'Table Grid'
headers = ['Test Category', 'Tests Executed', 'Tests Passed', 'Pass Rate']
for i, h in enumerate(headers):
    cell = table.rows[0].cells[i]; cell.text = h
    for p in cell.paragraphs:
        for r in p.runs: sf(r, 11, True)
data = [
    ['Functional Testing', '24', '24', '100%'],
    ['Structural Testing', '18', '17', '94.4%'],
    ['Integration Testing', '12', '12', '100%'],
    ['System Testing', '8', '8', '100%'],
    ['User Acceptance', '6', '6', '100%'],
]
for ri, row in enumerate(data):
    for ci, val in enumerate(row):
        cell = table.rows[ri+1].cells[ci]; cell.text = val
        for p in cell.paragraphs:
            for r in p.runs: sf(r, 11)

pb()

# ===== CHAPTER 8 =====
h1('CHAPTER 8: IMPLEMENTATION')
h2('8.1 Implementation of the Project')
body('Implementation followed the six-phase plan defined in Chapter 4. Each phase delivered independently verifiable components that were integrated progressively into the complete system.')
body('The Next.js application was initialized using the create-next-app CLI with the TypeScript App Router template. The Prisma schema was defined first, establishing the Student, Internship, Application, and Interaction models and their relationships before feature development began.')
body('The AI microservice was developed in Python using FastAPI, exposing three primary endpoints: POST /embed for generating BERT embeddings from text input; POST /recommend for executing the full SmartMatch-AI pipeline; and POST /update-weights for processing reinforcement learning feedback and adjusting the \u03b1 parameter.')
body('The blockchain verification module was integrated as a separate service that validates employer credentials and internship posting integrity through cryptographic hashing.')
body('The skill-gap analysis module compares the skills embedded in a student\u2019s profile vector against the competency requirements extracted from target internship descriptions.')

h2('8.2 Technology Stack Summary')
table2 = doc.add_table(rows=9, cols=2)
table2.style = 'Table Grid'
for i, h in enumerate(['Component', 'Technology']):
    cell = table2.rows[0].cells[i]; cell.text = h
    for p in cell.paragraphs:
        for r in p.runs: sf(r, 11, True)
stack = [
    ['Frontend', 'Next.js 15, TypeScript, Tailwind CSS'],
    ['Backend', 'Next.js API Routes, Python FastAPI'],
    ['Database', 'PostgreSQL (NeonDB Serverless)'],
    ['ORM', 'Prisma'],
    ['AI/ML', 'BERT (Hugging Face), scikit-learn, PyTorch'],
    ['Authentication', 'Clerk'],
    ['Deployment', 'Vercel'],
    ['APIs', 'Google Gemini API'],
]
for ri, row in enumerate(stack):
    for ci, val in enumerate(row):
        cell = table2.rows[ri+1].cells[ci]; cell.text = val
        for p in cell.paragraphs:
            for r in p.runs: sf(r, 11)

h2('8.3 Conversion Plan')
body('InternAI Compass was developed as a new system with no legacy software requiring migration in the context of this project. For institutional deployment, a migration plan would involve: exporting existing student and internship data from legacy systems in CSV format; running one-time Prisma migration scripts to populate the database; training placement officers through guided walkthroughs of the administrator dashboard; and running both systems in parallel for two to four weeks to validate output parity before full transition.')

h2('8.4 Post-Implementation and Software Maintenance')
body('The deployed system was monitored for three weeks following release to identify latency anomalies, authentication edge cases, and database connection pooling issues. No critical bugs were detected during this period.')
body('Ongoing maintenance for production deployment would include: periodic retraining of the BERT embedding model to incorporate emerging skill terminology; recalibration of the reinforcement learning learning rate \u03b7 based on observed convergence patterns; Prisma schema migrations to support new features; and renewal of API keys and service tokens prior to expiration.')

pb()

# ===== CHAPTER 9 =====
h1('CHAPTER 9: PROJECT LEGACY')
h2('9.1 Current Status of the Project')
body('All components of InternAI Compass have been successfully deployed, tested, and verified against the requirements specification. The live application is accessible at https://capstone-seminar.vercel.app/ and the complete source code is available at https://github.com/Mohan2347/capstone_seminar.')
body('The SmartMatch-AI algorithm achieved a recommendation accuracy of 93%, precision of 91%, and adaptability of 90% against a curated 100-internship expert benchmark, outperforming all compared baseline systems. The reinforcement learning module demonstrated convergence to optimal \u03b1 values within 50 interaction cycles in simulated testing. The blockchain verification module successfully authenticated all test employer registrations without false positives. The skill-gap analysis module correctly identified target competency deficiencies in 89% of test profile comparisons.')

h2('9.2 Remaining Areas of Concern')
body('Embedding Model Versioning: The current implementation uses a pre-trained BERT model that does not automatically update to incorporate emerging skills and terminology.')
body('Cold Start Problem: New students with minimal interaction history receive recommendations based primarily on content similarity until sufficient feedback is accumulated for effective collaborative filtering.')
body('Mobile Accessibility: The current interface targets desktop browsers. A native mobile application would significantly expand accessibility.')
body('Blockchain Scalability: The current blockchain verification approach is functional for demonstration scale but would require evaluation for throughput requirements at institutional deployment scale.')
body('Asynchronous Processing: Embedding generation is currently synchronous. For high concurrent usage, an asynchronous job queue would prevent latency spikes during peak periods.')

h2('9.3 Technical and Managerial Lessons Learnt')
body('Technical Lessons:')
body('Pre-computing internship embeddings at posting time rather than during recommendation generation reduced query latency by approximately 60%, demonstrating the value of anticipatory computation for performance-sensitive pipelines.')
body('Structured prompt engineering for the Gemini API integration required more iterations than anticipated. Specifying output format constraints explicitly in the prompt reduced parsing errors significantly.')
body('The schema-first approach to database design, where models and relationships were fully defined before feature development began, prevented the integration conflicts that commonly arise in multi-developer projects.')
body('Managerial Lessons:')
body('Dividing the project into six phases with explicit deliverables enabled clear progress tracking across a team of six developers. Aligning on shared TypeScript types before branching into parallel feature development was particularly effective.')
body('Embedding testing within the development workflow rather than reserving it for a final phase caught several edge cases earlier, reducing the effort required during the testing phase.')

pb()

doc.save(OUTPUT)
print('Part 3 done - chapters 5-9 added')
