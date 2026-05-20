import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def load_text(path):
    with open(path, 'r', encoding='utf-8') as f:
        return f.read()

def build_report():
    doc = Document()
    
    # Set default styles
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)
    font.color.rgb = RGBColor(0, 0, 0)

    # Helper for centered bold text
    def add_centered_bold(text, size=14):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(size)
        return p

    # --- TITLE PAGE ---
    add_centered_bold("CAPSTONE PROJECT REPORT", 16)
    add_centered_bold("(Project Term: January – May 2026)", 12)
    doc.add_paragraph("\n\n")
    add_centered_bold("INTERNAI COMPASS:", 18)
    add_centered_bold("AN INTELLIGENT HYBRID RECOMMENDATION ENGINE FOR PERSONALIZED AND EQUITABLE INTERNSHIP PATHWAYS", 14)
    doc.add_paragraph("\n\n")
    add_centered_bold("Submitted by", 12)
    
    students = [
        ("Y.N.V. Mahesh Reddy", "12201492"),
        ("T. Bhanu Prakesh", "12213376"),
        ("G.M.V.R. Reddy", "12206674"),
        ("Vemuri Gowtham", "12213211"),
        ("Dudekula Rahim", "12201812"),
        ("K. Venkata Sai", "12205733")
    ]
    for name, reg in students:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(f"{name}        Registration Number: {reg}")
    
    doc.add_paragraph("\n")
    add_centered_bold("Project Group Number: 2RGC0535", 12)
    add_centered_bold("Course Code: CSE439", 12)
    doc.add_paragraph("\n")
    add_centered_bold("Under the Guidance of", 12)
    add_centered_bold("Ms. Isha Khughar", 12)
    add_centered_bold("Assistant Professor", 11)
    doc.add_paragraph("\n\n")
    add_centered_bold("School of Computer Science and Engineering", 14)
    add_centered_bold("Lovely Professional University, Phagwara", 14)
    
    doc.add_page_break()

    # --- PAC FORM ---
    add_centered_bold("PAC FORM", 16)
    doc.add_paragraph("(To be attached as per university format)").alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()

    # --- DECLARATION ---
    add_centered_bold("DECLARATION", 16)
    p = doc.add_paragraph()
    p.add_run("We hereby declare that the project work entitled “InternAI Compass: An Intelligent Hybrid Recommendation Engine for Personalized and Equitable Internship Pathways” is an authentic record of our own work carried out as requirements of Capstone Project for the award of B.Tech degree in Computer Science and Engineering from Lovely Professional University, Phagwara, under the guidance of Ms. Isha Khughar, during January to May 2026. All the information furnished in this capstone project report is based on our own intensive work and is genuine.")
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    
    doc.add_paragraph("\nProject Group Number: 2RGC0535")
    for name, reg in students:
        doc.add_paragraph(f"Name of Student: {name}\nRegistration Number: {reg}\n(Signature of Student)")
    
    doc.add_page_break()

    # --- CERTIFICATE ---
    add_centered_bold("CERTIFICATE", 16)
    p = doc.add_paragraph()
    p.add_run("This is to certify that the declaration statement made by this group of students is correct to the best of my knowledge and belief. They have completed this Capstone Project under my guidance and supervision. The present work is the result of their original investigation, effort and study. No part of the work has ever been submitted for any other degree at any University. The Capstone Project is fit for the submission and partial fulfillment of the conditions for the award of B.Tech degree in Computer Science and Engineering from Lovely Professional University, Phagwara.")
    doc.add_paragraph("\n\n\nSignature and Name of the Mentor\nDesignation\nSchool of Computer Science and Engineering,\nLovely Professional University,\nPhagwara, Punjab.")
    doc.add_page_break()

    # --- TABLE OF CONTENTS ---
    add_centered_bold("TABLE OF CONTENTS", 16)
    toc_text = """CHAPTER 1: INTRODUCTION	9
CHAPTER 2: PROFILE OF THE PROBLEM, RATIONALE AND SCOPE	14
CHAPTER 3: EXISTING SYSTEM	16
CHAPTER 4: PROBLEM ANALYSIS	18
CHAPTER 5: SOFTWARE REQUIREMENT ANALYSIS	20
CHAPTER 6: DESIGN	23
CHAPTER 7: TESTING	28
CHAPTER 8: IMPLEMENTATION	31
CHAPTER 9: PROJECT LEGACY	34
CHAPTER 10: USER MANUAL	37
CHAPTER 11: SOURCE CODE AND SYSTEM SNAPSHOTS	39
CHAPTER 12: BIBLIOGRAPHY	61"""
    doc.add_paragraph(toc_text)
    doc.add_page_break()

    # --- CHAPTERS CONTENT ---
    chapters = [
        ("CHAPTER 1: INTRODUCTION", "Chapter_1_3.txt", "## CHAPTER 1:"),
        ("CHAPTER 2: PROFILE OF THE PROBLEM, RATIONALE AND SCOPE", "Chapter_1_3.txt", "## CHAPTER 2:"),
        ("CHAPTER 3: EXISTING SYSTEM", "Chapter_1_3.txt", "## CHAPTER 3:"),
        ("CHAPTER 4: PROBLEM ANALYSIS", "Chapter_4_6.txt", "## CHAPTER 4:"),
        ("CHAPTER 5: SOFTWARE REQUIREMENT ANALYSIS", "Chapter_4_6.txt", "## CHAPTER 5:"),
        ("CHAPTER 6: DESIGN", "Chapter_4_6.txt", "## CHAPTER 6:"),
        ("CHAPTER 7: TESTING", "Chapter_7_9.txt", "## CHAPTER 7:"),
        ("CHAPTER 8: IMPLEMENTATION", "Chapter_7_9.txt", "## CHAPTER 8:"),
        ("CHAPTER 9: PROJECT LEGACY", "Chapter_7_9.txt", "## CHAPTER 9:"),
        ("CHAPTER 10: USER MANUAL", "Chapter_10_12.txt", "## CHAPTER 10:"),
        ("CHAPTER 11: SOURCE CODE AND SYSTEM SNAPSHOTS", "Chapter_10_12.txt", "## CHAPTER 11:"),
        ("CHAPTER 12: BIBLIOGRAPHY", "Chapter_10_12.txt", "## CHAPTER 12:")
    ]

    img_dir = r"C:\Users\HP\.gemini\antigravity\brain\4d446f91-1ec0-4bc3-8feb-b0cd1930a8ce"
    images = {
        "ARCH": os.path.join(img_dir, "internai_compass_architecture_1777382356632.png"),
        "FLOW": os.path.join(img_dir, "smartmatch_ai_flowchart_1777382377701.png"),
        "LAND": os.path.join(img_dir, "landing_1777382830873.png"),
        "AUTH": os.path.join(img_dir, "auth_1777382859992.png"),
        "DASH": os.path.join(img_dir, "dashboard_1777382894561.png")
    }

    for title, filename, marker in chapters:
        add_centered_bold(title, 14)
        raw_content = load_text(filename)
        
        # Extract section text
        try:
            section_text = raw_content.split(marker)[1]
            # If not last chapter in that file, split at next header
            if "## CHAPTER" in section_text:
                section_text = section_text.split("## CHAPTER")[0]
        except:
            section_text = raw_content

        # Add paragraphs
        for line in section_text.split('\n'):
            if line.strip():
                # Handle Subheadings
                if line.strip().startswith('###') or line.strip().isupper() or (len(line.strip()) < 50 and '.' in line[:5]):
                    p = doc.add_paragraph()
                    run = p.add_run(line.replace('###', '').strip())
                    run.bold = True
                else:
                    p = doc.add_paragraph(line.strip())
                    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

        # Special Case: Images
        if "CHAPTER 6" in title:
            # Insert Arch after System Design
            doc.add_paragraph("\n")
            if os.path.exists(images["ARCH"]):
                doc.add_picture(images["ARCH"], width=Inches(5.5))
                doc.add_paragraph("Figure 6.1: InternAI Compass System Architecture").alignment = WD_ALIGN_PARAGRAPH.CENTER
            
            doc.add_paragraph("\n")
            if os.path.exists(images["FLOW"]):
                doc.add_picture(images["FLOW"], width=Inches(5.5))
                doc.add_paragraph("Figure 6.2: SmartMatch-AI Algorithm Flowchart").alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        if "CHAPTER 11" in title:
            if os.path.exists(images["LAND"]):
                doc.add_picture(images["LAND"], width=Inches(5.5))
                doc.add_paragraph("Figure 11.1: Project Landing Page").alignment = WD_ALIGN_PARAGRAPH.CENTER
            if os.path.exists(images["AUTH"]):
                doc.add_picture(images["AUTH"], width=Inches(5.5))
                doc.add_paragraph("Figure 11.2: Authentication Interface").alignment = WD_ALIGN_PARAGRAPH.CENTER
            if os.path.exists(images["DASH"]):
                doc.add_picture(images["DASH"], width=Inches(5.5))
                doc.add_paragraph("Figure 11.3: Internship Discovery Dashboard").alignment = WD_ALIGN_PARAGRAPH.CENTER

        doc.add_page_break()

    doc.save(r"C:\Users\HP\Desktop\COMPLETE_INTERNAI_COMPASS_FINAL_REPORT.docx")
    print("Done")

if __name__ == "__main__":
    build_report()
