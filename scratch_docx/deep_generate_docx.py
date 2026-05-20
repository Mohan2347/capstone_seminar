import os
from docx import Document
from docx.shared import Pt

# Load the comprehensive content we generated for InternAI Compass
content_file = r"C:\Users\HP\Desktop\internship-recommendation\scratch_docx\InternAI_Compass_Final_Report.md"
with open(content_file, 'r', encoding='utf-8') as f:
    full_content = f.read()

def get_section(chapter_num):
    marker = f"## CHAPTER {chapter_num}:"
    next_marker = f"## CHAPTER {chapter_num + 1}:"
    try:
        start = full_content.find(marker)
        if start == -1: return ""
        end = full_content.find(next_marker)
        if end == -1:
            return full_content[start:].strip()
        return full_content[start:end].strip()
    except:
        return ""

def deep_process_docx():
    template_path = r"C:\Users\HP\Downloads\ganesh_report.docx"
    output_path = r"C:\Users\HP\Desktop\InternAI_Compass_Final_Report_V2.docx"
    
    doc = Document(template_path)
    
    # Define primary identity replacements
    replacements = {
        "AI-EMPOWERED URBAN GOVERNANCE SYSTEM": "INTERNAI COMPASS",
        "CIVIC ISSUE REPORTING": "INTERNSHIP RECOMMENDATION",
        "citizen": "student",
        "Citizen": "Student",
        "administrator": "admin/employer",
        "Guna444/AI-Based-Civic-Issue-Reporting-System": "Mohan2347/capstone_seminar",
        "ai-civic-issue.vercel.app": "capstone-seminar.vercel.app",
        "https://ai-civic-issue.vercel.app/": "https://capstone-seminar.vercel.app/",
        "MCIA": "SMARTMATCH-AI"
    }

    # 1. Identity Replacement across all paragraphs
    for p in doc.paragraphs:
        for key, val in replacements.items():
            if key in p.text:
                p.text = p.text.replace(key, val)

    # 2. Student List Replacement (Hardcoded from your files)
    student_replacements = {
        "Nithesh Reddy Minam Reddy Gari": "Y.N.V. Mahesh Reddy",
        "Krishnan Guna Shaker": "T. Bhanu Prakesh",
        "S Deeraj Kumar Reddy": "G.M.V.R. Reddy",
        "Ganesh Maddala": "Vemuri Gowtham",
        "M Siva Nagi Reddy": "Dudekula Rahim",
        "Venkata Naga Chiranjeevi Kotari": "K. Venkata Sai"
    }
    for p in doc.paragraphs:
        for old, new in student_replacements.items():
            if old in p.text:
                p.text = p.text.replace(old, new)

    # 3. Chapter Content Replacement
    # We will look for chapter headers and replace the paragraphs immediately following them.
    # Note: This is an invasive edit to ensure the friend's content is gone.
    
    for i in range(1, 13):
        new_text = get_section(i)
        if not new_text: continue
        
        # Strip the markdown header from the text
        clean_text = new_text.replace(f"## CHAPTER {i}:", "").strip()
        
        # We search for the chapter heading in the doc
        for idx, p in enumerate(doc.paragraphs):
            if f"CHAPTER {i}:" in p.text.upper():
                # We found the chapter start. We'll replace the next few paragraphs 
                # until we hit the next CHAPTER heading.
                # To be safe and preserve "pages", we'll just overwrite the text of the found paragraph's followers.
                next_idx = idx + 1
                paragraphs_overwritten = 0
                while next_idx < len(doc.paragraphs) and "CHAPTER" not in doc.paragraphs[next_idx].text.upper():
                    # We clear the existing text and only write the new text into the first paragraph after heading
                    if paragraphs_overwritten == 0:
                        doc.paragraphs[next_idx].text = clean_text
                    else:
                        doc.paragraphs[next_idx].text = "" # Clear remaining paragraphs in this section
                    next_idx += 1
                    paragraphs_overwritten += 1

    doc.save(output_path)
    print(f"Deeply processed report saved to {output_path}")

if __name__ == "__main__":
    deep_process_docx()
